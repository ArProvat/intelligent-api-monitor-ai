from pathlib import Path
from docx import Document

source = Path('docs/ai-prompts.md')
document = Document()
in_code = False
for line in source.read_text(encoding='utf-8').splitlines():
    if line.startswith(('```', '~~~')):
        in_code = not in_code
    elif line.startswith('# ') and not in_code:
        document.add_heading(line[2:], level=1)
    elif line.startswith('## ') and not in_code:
        document.add_heading(line[3:], level=2)
    else:
        paragraph = document.add_paragraph(line)
        if in_code:
            for run in paragraph.runs:
                run.font.name = 'Consolas'
document.save('AI_Prompts.docx')
